#!/usr/bin/env python3
"""Export visible Excel worksheets to UTF-8 CSV without requiring Excel."""

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_PATH = Path(__file__).resolve().parent.parent / "_easyConfig_import"


def find_workbooks(path: Path, recurse: bool) -> list[Path]:
    if not path.exists():
        raise FileNotFoundError(f"Input path does not exist: {path}")
    if path.is_file():
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Input must be an .xlsx workbook.")
        return [path]
    candidates = path.rglob("*") if recurse else path.iterdir()
    return sorted(
        file
        for file in candidates
        if file.is_file()
        and file.suffix.lower() == ".xlsx"
        and not file.name.startswith("~$")
    )


def export_workbook(path: Path, dry_run: bool = False) -> int:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [
            sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"
        ]
        if not sheets:
            raise ValueError(f"No visible worksheets in {path}")
        for sheet in sheets:
            name = (
                path.stem
                if len(sheets) == 1
                else f"{path.stem}_{sheet.title}"
            )
            destination = path.with_name(f"{name}.csv")
            if dry_run:
                print(f"Would create {destination} (sheet: {sheet.title})")
                continue
            with destination.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerows(sheet.iter_rows(values_only=True))
            print(f"Created {destination} (sheet: {sheet.title})")
        return len(sheets)
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "path",
        nargs="?",
        type=Path,
        default=DEFAULT_PATH,
        help="workbook or folder (default: repository's _easyConfig_import folder)",
    )
    parser.add_argument(
        "--recurse", action="store_true", help="search subfolders for workbooks"
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="preview exports without writing CSVs"
    )
    args = parser.parse_args()
    workbooks = find_workbooks(args.path, args.recurse)
    if not workbooks:
        parser.error(f"No .xlsx workbooks found under {args.path}")
    count = sum(export_workbook(path, args.dry_run) for path in workbooks)
    print(f"Done. CSV files {'planned' if args.dry_run else 'written'}: {count}")


if __name__ == "__main__":
    main()
